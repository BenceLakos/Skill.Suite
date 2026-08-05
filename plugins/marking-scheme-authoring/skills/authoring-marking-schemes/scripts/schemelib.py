"""Shared reader for a CIS marking scheme workbook.

Every rule about the sheet's shape lives here once — which rows are aspect rows, where the
sub-criterion labels sit, which column is the expected result and which the input. Five copies of
"an aspect row is one whose column D is M or J" is exactly the drift this toolchain exists to stop.

See references/cis-format.md for the full anatomy.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal

try:
    import openpyxl
except ImportError:  # pragma: no cover - dependency guidance is more useful than a stack trace
    raise SystemExit("openpyxl is required:  pip3 install openpyxl")

SHEET_CIS = "CIS Marking Scheme Import"
SHEET_TEST_MAP = "Test Map"
SHEET_CALCULATIONS = "Calculations"

ASPECT_ID = re.compile(r"^[A-Z]\d+\.\d+$")
SUBCRIT_ID = re.compile(r"^[A-Z]\d+$")

# Column indices (1-based) of an aspect row on the CIS sheet.
COL_SUBCRIT_ID, COL_SUBCRIT_NAME, COL_DAY = 1, 2, 3
COL_TYPE, COL_DESCRIPTION, COL_JUDG_SCORE = 4, 5, 6
COL_EXPECTED, COL_REQUIREMENT, COL_WSOS, COL_CALC_ROW, COL_MARK = 7, 8, 9, 10, 11


def _dec(value) -> Decimal:
    """Marks as Decimal. Float summation of a mark ladder is not guaranteed exact."""
    return Decimal(str(value)) if value is not None else Decimal(0)


def _text(value) -> str:
    return "" if value is None else str(value).strip()


@dataclass
class Aspect:
    """One graded aspect. `id` comes from the Test Map — the CIS sheet does not carry it."""

    id: str
    row: int
    criterion: str
    subcrit: str
    subcrit_name: str
    kind: str          # "M" measurement or "J" judgement
    description: str   # CIS column E, the marking line
    expected: str      # CIS column G — what the implementation must produce
    requirement: str   # CIS column H — what it is given
    wsos: int
    mark: Decimal
    test: str = ""     # Test Map column F; "" for a pipeline-judged aspect

    @property
    def is_pipeline_judged(self) -> bool:
        return self.test == "" or self.test.startswith("(")


@dataclass
class Scheme:
    path: str
    aspects: list[Aspect] = field(default_factory=list)
    declared_wsos: dict[int, Decimal] = field(default_factory=dict)
    declared_criteria: dict[str, Decimal] = field(default_factory=dict)
    # Criterion id -> the formula in its Criteria-table mark cell, when it holds one instead of a
    # literal. A generated workbook has no cached values, so this is the only evidence there.
    criteria_formulas: dict[str, str] = field(default_factory=dict)
    declared_total: Decimal = Decimal(0)
    block_sum_ranges: dict[str, tuple[int, int]] = field(default_factory=dict)
    block_total_cells: dict[str, str] = field(default_factory=dict)
    sumif_range: tuple[int, int] | None = None
    test_map: dict[str, str] = field(default_factory=dict)
    test_map_lines: dict[str, str] = field(default_factory=dict)

    def total(self) -> Decimal:
        return sum((a.mark for a in self.aspects), Decimal(0))

    def by_criterion(self) -> dict[str, Decimal]:
        out: dict[str, Decimal] = {}
        for a in self.aspects:
            out[a.criterion] = out.get(a.criterion, Decimal(0)) + a.mark
        return out

    def by_wsos(self) -> dict[int, Decimal]:
        out: dict[int, Decimal] = {}
        for a in self.aspects:
            out[a.wsos] = out.get(a.wsos, Decimal(0)) + a.mark
        return out

    def graded(self) -> list[Aspect]:
        return [a for a in self.aspects if not a.is_pipeline_judged]


def read_scheme(path: str) -> Scheme:
    """Parse a scheme workbook. Reads formulas AND cached values: the reconciliation cells are
    formulas, and a workbook openpyxl wrote has no cached values for them at all."""
    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)

    for name in (SHEET_CIS, SHEET_TEST_MAP):
        if name not in values.sheetnames:
            raise SystemExit(f"{path}: missing sheet {name!r} (found {values.sheetnames})")

    scheme = Scheme(path=path)
    _read_test_map(values[SHEET_TEST_MAP], scheme)
    _read_cis(values[SHEET_CIS], formulas[SHEET_CIS], scheme)
    return scheme


def _read_test_map(ws, scheme: Scheme) -> None:
    order: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cells = list(row) + [None] * 8
        aspect_id = _text(cells[0])
        if not ASPECT_ID.match(aspect_id):
            continue
        order.append(aspect_id)
        scheme.test_map[aspect_id] = _text(cells[5])
        scheme.test_map_lines[aspect_id] = _text(cells[4])
    scheme._order = order  # type: ignore[attr-defined]


def _read_cis(ws, wsf, scheme: Scheme) -> None:
    # Declared WSOS marks: column A the section number, I the declared marks.
    for r in range(1, ws.max_row + 1):
        section = ws.cell(r, 1).value
        if isinstance(section, int) and 1 <= section <= 6 and ws.cell(r, 9).value is not None:
            scheme.declared_wsos[section] = _dec(ws.cell(r, 9).value)

    # Criteria table: a single-letter id in A, the mark in K. Sits above the first block header.
    # The mark may be a literal (an Excel-saved file caches it) or a formula pointing at the block
    # total (a generated file has no cache at all) — record whichever is there.
    for r in range(1, ws.max_row + 1):
        cid = _text(ws.cell(r, 1).value)
        mark = ws.cell(r, 11).value
        formula = wsf.cell(r, 11).value
        is_criterion = len(cid) == 1 and cid.isalpha() and cid.isupper()
        if is_criterion and isinstance(mark, (int, float)):
            scheme.declared_criteria[cid] = _dec(mark)
        elif is_criterion and isinstance(formula, str) and formula.startswith("="):
            scheme.criteria_formulas[cid] = formula.replace(" ", "")
        if cid.lower().startswith("session total") or _text(ws.cell(r, 2).value).lower() == "session total":
            if isinstance(mark, (int, float)):
                scheme.declared_total = _dec(mark)

    # Formula-derived ranges, for the "does the declared range match the real block?" check.
    for r in range(1, wsf.max_row + 1):
        for c in range(1, min(wsf.max_column, 20) + 1):
            v = wsf.cell(r, c).value
            if not isinstance(v, str) or not v.startswith("="):
                continue
            m = re.match(r"=SUM\(K(\d+):K(\d+)\)$", v.replace(" ", ""))
            if m:
                # Only a criterion block total, identified by the "Criterion X" label two cells left.
                # The WSOS table's own =SUM(K5:K10) total-variation cell has the same shape.
                label = _text(wsf.cell(r, c - 2).value)
                if label.lower().startswith("criterion"):
                    scheme.block_sum_ranges[label] = (int(m.group(1)), int(m.group(2)))
                    scheme.block_total_cells[label] = f"{wsf.cell(r, c).column_letter}{r}"
            m = re.search(r"SUMIF\(\$I\$(\d+):\$I\$(\d+)", v.replace(" ", ""))
            if m and scheme.sumif_range is None:
                scheme.sumif_range = (int(m.group(1)), int(m.group(2)))

    # Aspect rows, in sheet order. Ids are positional — zip against the Test Map's order.
    order: list[str] = getattr(scheme, "_order", [])
    criterion = ""
    subcrit = subcrit_name = ""
    index = 0
    for r in range(1, ws.max_row + 1):
        a_cell = _text(ws.cell(r, COL_SUBCRIT_ID).value)
        if SUBCRIT_ID.match(a_cell):
            subcrit, criterion = a_cell, a_cell[0]
            subcrit_name = _text(ws.cell(r, COL_SUBCRIT_NAME).value)
            continue
        kind = _text(ws.cell(r, COL_TYPE).value).upper()
        if kind not in ("M", "J"):
            continue
        aspect_id = order[index] if index < len(order) else f"?{index + 1}"
        index += 1
        wsos_raw = ws.cell(r, COL_WSOS).value
        scheme.aspects.append(
            Aspect(
                id=aspect_id,
                row=r,
                criterion=criterion,
                subcrit=subcrit,
                subcrit_name=subcrit_name,
                kind=kind,
                description=_text(ws.cell(r, COL_DESCRIPTION).value),
                expected=_text(ws.cell(r, COL_EXPECTED).value),
                requirement=_text(ws.cell(r, COL_REQUIREMENT).value),
                wsos=int(wsos_raw) if isinstance(wsos_raw, (int, float)) else 0,
                mark=_dec(ws.cell(r, COL_MARK).value),
                test=scheme.test_map.get(aspect_id, ""),
            )
        )


class Findings:
    """Collects errors and warnings; exit code comes from whether any error was recorded."""

    def __init__(self, title: str) -> None:
        self.title = title
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def error(self, message: str) -> None:
        self.errors.append(message)

    def warn(self, message: str) -> None:
        self.warnings.append(message)

    def report(self) -> int:
        print(self.title)
        for w in self.warnings:
            print(f"  warn  {w}")
        for e in self.errors:
            print(f"  FAIL  {e}")
        if not self.errors and not self.warnings:
            print("  ok")
        elif not self.errors:
            print(f"  ok, with {len(self.warnings)} warning(s)")
        else:
            print(f"  {len(self.errors)} error(s), {len(self.warnings)} warning(s)")
        return 1 if self.errors else 0
